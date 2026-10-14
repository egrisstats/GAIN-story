"""Fill the untranslated French, Spanish, Russian and Arabic cells of the GAIN 2026 Kobo form.

Usage: python3 kobo_fill.py <translations dir> <input arm .xlsx> <output .xlsx>

<translations dir> holds todo.json (the English strings, with ids) and out_<lang>.json (one
translation per id). A cell is filled when its English text is one of the strings in todo.json and its translation is
empty, a copy of the English (in any case), or incomplete: the English quotes an answer option
("Don't know") and the existing translation quotes nothing, so its second sentence is missing.
Nothing else in the form changes.
Name lists (countries, institutions) are left alone.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

LANGS = {"fr": "French (fr)", "es": "Spanish (es)", "ru": "Russian (ru)", "ar": "Arabic (ar)"}
NAME_LISTS = {"Country", "Institution", "Bureau", "LOC_academia_list", "LOC_cso_list", "Region"}


def plain(v):
    return re.sub(r"<[^>]+>|\$\{[^}]+\}", "", str(v)).strip().lower()


def incomplete(cur, en):
    return '"' in str(en) and not re.search(r'["«»“”„]', str(cur))


def main(tdir, src, dst):
    tdir = Path(tdir)
    todo = json.loads((tdir / "todo.json").read_text(encoding="utf-8"))
    by_en = {t["en"]: t["id"] for t in todo}
    tr = {c: {o["id"]: o[c] for o in json.loads((tdir / f"out_{c}.json").read_text(encoding="utf-8"))} for c in LANGS}

    wb = openpyxl.load_workbook(src)
    filled = []
    for sheet in ("survey", "choices"):
        ws = wb[sheet]
        head = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(head) if h}
        for row in ws.iter_rows(min_row=2):
            get = lambda h: row[col[h]].value if h in col else None
            if sheet == "choices" and get("list_name") in NAME_LISTS:
                continue
            for h in head:
                if not h or not str(h).endswith("::English (en)") or h.startswith("media"):
                    continue
                en = get(h)
                if en is None or str(en) not in by_en:
                    continue
                base = h.split("::")[0]
                for code, lang in LANGS.items():
                    target = f"{base}::{lang}"
                    if target not in col:
                        continue
                    cur = get(target)
                    if cur and plain(cur) != plain(en) and not incomplete(cur, en):
                        continue
                    new = tr[code][by_en[str(en)]]
                    if cur and str(cur).isupper():  # older choice lists are written in capitals
                        new = new.upper()
                    row[col[target]].value = new
                    filled.append((sheet, get("list_name") or "", get("name"), base, code))
    wb.save(dst)
    print(f"{dst}: filled {len(filled)} cells")
    return filled


if __name__ == "__main__":
    main(*sys.argv[1:4])
